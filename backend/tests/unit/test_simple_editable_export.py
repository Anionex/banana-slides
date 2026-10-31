"""Text-only export keeps artwork in the background and never doubles failed text."""
import io
import threading
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import Mock

import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from models import Page, Project, Task, db
from services.export_service import ExportError, ExportService
from services.image_editability import ServiceConfig, ImageEditabilityService
from services.image_editability.data_models import BBox, EditableElement, EditableImage
from services.image_editability.extractors import BaiduAccurateOCRElementExtractor
from services.image_editability.factories import ExtractorFactory, InpaintProviderFactory
from services.image_editability.inpaint_providers import GenerativeEditInpaintProvider
from services.prompts import get_clean_background_prompt


@pytest.fixture
def simple_services(tmp_path, monkeypatch):
    source = tmp_path / 'slide.png'
    image = Image.new('RGB', (320, 180), 'white')
    draw = ImageDraw.Draw(image)
    draw.rectangle((200, 80, 280, 150), fill='blue')
    draw.text((20, 20), 'Editable text', fill='black')
    draw.text((20, 80), 'Missed by OCR', fill='black')
    image.save(source)

    ocr = Mock()
    ocr.recognize.return_value = {
        'image_size': image.size,
        'text_lines': [{'text': 'Editable text', 'bbox': [20, 20, 120, 40]}],
    }
    original_ocr = ocr.recognize.return_value
    ocr.recognize.side_effect = lambda image_path, **kwargs: (original_ocr if image_path == str(source) else {'image_size': image.size, 'text_lines': []})
    inpaint = Mock()
    background = Image.new('RGB', image.size, 'white')
    ImageDraw.Draw(background).rectangle((200, 80, 280, 150), fill='blue')
    inpaint.inpaint_regions.return_value = background
    monkeypatch.setattr(ExtractorFactory, 'create_baidu_accurate_ocr_extractor',
                        lambda: BaiduAccurateOCRElementExtractor(ocr))
    monkeypatch.setattr(InpaintProviderFactory, 'create_generative_edit_provider',
                        lambda **_kwargs: inpaint)
    return source, ocr, inpaint


def build_simple(simple_services, tmp_path):
    source, _, _ = simple_services
    config = ServiceConfig.from_defaults(
        upload_folder=str(tmp_path), text_only=True, inpaint_method='hybrid',
        max_depth=5, enable_icon_subject_extraction=True,
    )
    assert config.max_depth == 1
    assert config.enable_icon_subject_extraction is False
    return ImageEditabilityService(config).make_image_editable(str(source))


def export_file(tmp_path, **options):
    output = tmp_path / 'result.pptx'
    _, warnings = ExportService.create_editable_pptx_with_recursive_analysis(
        output_file=str(output), **options,
    )
    return output.read_bytes(), warnings


def test_simple_ocr_without_mineru_preserves_artwork_and_exports_text(simple_services, tmp_path):
    result = build_simple(simple_services, tmp_path)
    _, ocr, inpaint = simple_services
    assert ocr.recognize.call_count == 2
    assert inpaint.inpaint_regions.call_args.kwargs['types'] == ['text']
    assert inpaint.inpaint_regions.call_args.kwargs['text_only'] is True
    assert inpaint.inpaint_regions.call_args.kwargs['text_targets'] == [
        {'text': 'Editable text', 'bbox': [20, 20, 120, 40]},
    ]
    assert [elem.element_type for elem in result.elements] == ['text']
    with Image.open(result.clean_background) as background:
        assert background.getpixel((220, 100)) == (0, 0, 255)
        assert background.crop((20, 20, 120, 40)).getextrema() == ((255, 255),) * 3
        # The whole model output is used: OCR omissions must not be pasted back.
        assert background.crop((20, 80, 120, 100)).getextrema() == ((255, 255),) * 3
        assert background.tobytes() == inpaint.inpaint_regions.return_value.tobytes()

    pptx, warnings = export_file(tmp_path,
        editable_images=[result], slide_width_pixels=320, slide_height_pixels=180,
        export_mode='simple',
    )
    prs = Presentation(io.BytesIO(pptx))
    assert not warnings.has_warnings()
    assert len(prs.slides[0].shapes) == 2
    assert prs.slides[0].shapes[0].shape_type == MSO_SHAPE_TYPE.PICTURE
    assert prs.slides[0].shapes[1].text == 'Editable text'
    assert prs.slide_width / prs.slide_height == pytest.approx(16 / 9)
    text_shape = prs.slides[0].shapes[1]
    assert text_shape.left / prs.slide_width == pytest.approx(20 / 320, abs=.005)
    assert text_shape.top / prs.slide_height == pytest.approx(20 / 180, abs=.005)


@pytest.mark.parametrize('failure', ['none', 'exception', 'unchanged'])
@pytest.mark.parametrize('fail_fast', [True, False])
def test_failed_removal_never_overlays_original_text(simple_services, tmp_path, failure, fail_fast):
    source, _, inpaint = simple_services
    inpaint.inpaint_regions.return_value = None
    if failure == 'exception':
        inpaint.inpaint_regions.side_effect = RuntimeError('repair failed')
    elif failure == 'unchanged':
        inpaint.inpaint_regions.return_value = Image.open(source).copy()
    result = build_simple(simple_services, tmp_path)
    assert result.clean_background is None
    options = dict(editable_images=[result], export_mode='simple', fail_fast=fail_fast)
    if fail_fast:
        with pytest.raises(ExportError, match='文字擦除失败') as error:
            ExportService.create_editable_pptx_with_recursive_analysis(**options)
        assert error.value.error_code == 'EXPORT_TEXT_REMOVAL_FAILED'
    else:
        pptx, warnings = export_file(tmp_path, **options)
        shapes = Presentation(io.BytesIO(pptx)).slides[0].shapes
        assert len(shapes) == 1
        assert shapes[0].shape_type == MSO_SHAPE_TYPE.PICTURE
        assert any('该页文字不可编辑' in warning for warning in warnings.to_summary())


def test_no_text_keeps_original_without_inpainting(simple_services, tmp_path):
    _, ocr, inpaint = simple_services
    ocr.recognize.return_value['text_lines'] = []
    result = build_simple(simple_services, tmp_path)
    inpaint.inpaint_regions.assert_not_called()
    pptx, warnings = export_file(tmp_path,
        editable_images=[result], export_mode='simple',
    )
    assert len(Presentation(io.BytesIO(pptx)).slides[0].shapes) == 1
    assert not warnings.has_warnings()


@pytest.mark.parametrize('masked', [False, True])
@pytest.mark.parametrize('explicit', [False, True])
def test_targeted_erasure_uses_overlay_text_and_optional_alpha_mask(tmp_path, masked, explicit):
    import json
    ai = Mock()
    ai.image_provider.supports_mask_edit = masked
    source = Image.new('RGB', (320, 180), 'blue')
    ai.edit_image.return_value = source.copy()
    targets = [{'text': '① 删除指定文字', 'bbox': [20, 20, 120, 40]}]
    result = GenerativeEditInpaintProvider(ai).inpaint_regions(
        source, [targets[0]['bbox']], text_only=True, text_targets=targets,
        save_mask_path=str(tmp_path / 'mask.png'),
        **({'use_text_mask': masked} if explicit else {}),
    )
    assert result is not None
    request = ai.edit_image.call_args.kwargs
    assert '① 删除指定文字' in request['prompt']
    assert '[0.0625,0.11111,0.375,0.22222]' in request['prompt']
    assert ('mask_image' in request) is masked
    if masked:
        mask = request['mask_image']
        assert mask.size == source.size
        assert mask.getpixel((30, 30))[3] == 0
        assert mask.getpixel((200, 100))[3] == 255
    audit = json.loads((tmp_path / 'text_edit_request.json').read_text())
    assert audit['targets'] == targets
    assert audit['mask_supplied'] is masked


@pytest.mark.parametrize('failure', ['residual', 'ocr_error'])
@pytest.mark.parametrize('fail_fast', [True, False])
def test_unverified_background_never_overlays_text(simple_services, tmp_path, failure, fail_fast):
    _, ocr, _ = simple_services
    original = ocr.recognize.return_value
    ocr.recognize.side_effect = [original, original if failure == 'residual' else RuntimeError('OCR unavailable')]
    result = build_simple(simple_services, tmp_path)
    assert result.clean_background is None
    if fail_fast:
        with pytest.raises(ExportError, match='文字擦除失败'):
            export_file(tmp_path, editable_images=[result], export_mode='simple', fail_fast=True)
    else:
        data, warnings = export_file(tmp_path, editable_images=[result], export_mode='simple', fail_fast=False)
        shapes = Presentation(io.BytesIO(data)).slides[0].shapes
        assert len(shapes) == 1 and shapes[0].shape_type == MSO_SHAPE_TYPE.PICTURE
        assert warnings.has_warnings()


def test_preserved_icon_ocr_outside_text_boxes_does_not_block_export(simple_services, tmp_path):
    _, ocr, _ = simple_services
    original = ocr.recognize.return_value
    ocr.recognize.side_effect = [original, {'image_size': (320, 180), 'text_lines': [
        {'text': 'A', 'bbox': [210, 80, 240, 110]},
    ]}]
    result = build_simple(simple_services, tmp_path)
    assert result.clean_background is not None


def test_simple_requires_ocr_without_falling_back_to_layout(tmp_path, monkeypatch):
    monkeypatch.setattr(ExtractorFactory, 'create_baidu_accurate_ocr_extractor', lambda: None)
    with pytest.raises(ValueError, match='Baidu API Key'):
        ServiceConfig.from_defaults(text_only=True, upload_folder=str(tmp_path))


def test_simple_prompt_preserves_charts():
    prompt = get_clean_background_prompt([{'element_type': 'text'}], text_only=True)
    assert '保留所有图片、插画、图表' in prompt
    assert '移除整张图中的所有文字' in prompt
    assert '区域坐标' not in prompt
    assert '彻底移除页面中的所有文字、插画、图表' not in prompt
    assert '彻底移除页面中的所有文字、插画、图表' in get_clean_background_prompt([])


@pytest.mark.parametrize('method', [None, 'hybrid', 'baidu', 'generative'])
def test_simple_always_uses_image_model(simple_services, tmp_path, monkeypatch, method):
    baidu = Mock(side_effect=AssertionError('Simple mode must not create Baidu repair'))
    hybrid = Mock(side_effect=AssertionError('Simple mode must not create hybrid repair'))
    monkeypatch.setattr(InpaintProviderFactory, 'create_baidu_inpaint_provider', baidu)
    monkeypatch.setattr(InpaintProviderFactory, 'create_hybrid_inpaint_provider', hybrid)
    config = ServiceConfig.from_defaults(
        text_only=True, upload_folder=str(tmp_path), inpaint_method=method,
    )
    assert config.inpaint_registry.get_all_providers() == [simple_services[2]]
    baidu.assert_not_called()
    hybrid.assert_not_called()


@pytest.mark.parametrize('size,ratio', [((320, 180), '16:9'), ((240, 180), '4:3'), ((180, 320), '9:16')])
def test_simple_edit_uses_source_image_and_ratio(size, ratio):
    source = Image.new('RGB', size, 'blue')
    ai = Mock()

    def edit_image(**kwargs):
        with Image.open(kwargs['current_image_path']) as reference:
            assert reference.size == size
            assert reference.tobytes() == source.tobytes()
        assert kwargs['aspect_ratio'] == ratio
        assert '移除整张图中的所有文字' in kwargs['prompt']
        return source.copy()

    ai.edit_image.side_effect = edit_image
    result = GenerativeEditInpaintProvider(ai).inpaint_regions(
        source, [(10, 10, 50, 30)], types=['text'], text_only=True,
    )
    assert result.size == size
    ai.edit_image.assert_called_once()


@pytest.mark.parametrize('sizes', [[(320, 180), (240, 180)], [(240, 180), (320, 180)]])
def test_mixed_sizes_keep_text_aligned_with_scaled_background(tmp_path, sizes):
    pages = []
    for index, (width, height) in enumerate(sizes):
        source = tmp_path / f'{index}.png'
        Image.new('RGB', (width, height), 'white').save(source)
        bbox = BBox(width * .1, height * .2, width * .4, height * .3)
        pages.append(EditableImage(str(index), str(source), width, height,
            clean_background=str(source), elements=[
                EditableElement(str(index), 'text', bbox, bbox, content='Aligned'),
            ]))
    pptx, _ = export_file(tmp_path, editable_images=pages, export_mode='simple',
                          slide_width_pixels=sizes[0][0], slide_height_pixels=sizes[0][1])
    prs = Presentation(io.BytesIO(pptx))
    assert prs.slide_width / prs.slide_height == pytest.approx(sizes[0][0] / sizes[0][1])
    for slide, source_size in zip(prs.slides, sizes):
        background, text = slide.shapes
        assert background.image.size == source_size
        assert (background.width, background.height) == (prs.slide_width, prs.slide_height)
        assert text.left / background.width == pytest.approx(.1, abs=.005)
        assert text.top / background.height == pytest.approx(.2, abs=.005)
        assert text.width / background.width == pytest.approx(.3, abs=.005)
        assert text.height / background.height == pytest.approx(.1, abs=.005)


def test_simple_failure_does_not_wait_for_sibling_or_publish(tmp_path, monkeypatch):
    started, release = threading.Event(), threading.Event()
    bbox = BBox(0, 0, 10, 10)
    failed = EditableImage('bad', 'bad.png', 100, 100, elements=[
        EditableElement('text', 'text', bbox, bbox, content='failed'),
    ])
    monkeypatch.setattr(ServiceConfig, 'from_defaults', lambda **_kwargs: object())
    monkeypatch.setattr(ImageEditabilityService, '__init__', lambda *_args: None)

    def analyze(_self, path):
        if path == 'slow.png':
            started.set()
            assert release.wait(5)
            return EditableImage('slow', path, 100, 100)
        assert started.wait(3)
        return failed

    monkeypatch.setattr(ImageEditabilityService, 'make_image_editable', analyze)
    output = tmp_path / 'failed.pptx'
    with ThreadPoolExecutor(max_workers=1) as executor:
        future = executor.submit(
            ExportService.create_editable_pptx_with_recursive_analysis,
            image_paths=['slow.png', 'bad.png'], output_file=str(output),
            export_mode='simple', max_workers=2,
        )
        try:
            with pytest.raises(ExportError, match='文字擦除失败'):
                future.result(timeout=3)
            assert not output.exists()
        finally:
            release.set()
    assert not output.exists()


def test_simple_endpoint_task_download_and_selected_pages(app, client, simple_services, monkeypatch):
    source, ocr, _ = simple_services
    project = Project(creation_type='idea', idea_prompt='simple export', export_inpaint_method='baidu')
    db.session.add(project)
    db.session.flush()
    selected = Page(project_id=project.id, order_index=0, generated_image_path=str(source))
    db.session.add_all([selected, Page(project_id=project.id, order_index=1)])
    db.session.commit()
    project_id, page_id = project.id, selected.id
    monkeypatch.setattr('services.file_service.FileService.get_absolute_path', lambda *_args: str(source))
    monkeypatch.setattr('services.image_editability.TextAttributeExtractorFactory.create_caption_model_extractor', lambda: None)
    monkeypatch.setattr('services.task_manager.task_manager.submit_task',
                        lambda task_id, fn, **kwargs: fn(task_id, **kwargs))
    response = client.post(f'/api/projects/{project_id}/export/editable-pptx', json={
        'export_mode': 'simple', 'page_ids': [page_id],
    })
    assert response.status_code == 200
    task = db.session.get(Task, response.json['data']['task_id'])
    assert task.status == 'COMPLETED', task.error_message
    progress = task.get_progress()
    assert progress['export_mode'] == 'simple'
    assert db.session.get(Project, project_id).export_inpaint_method == 'baidu'
    assert ocr.recognize.call_count == 2
    download = client.get(progress['download_url'])
    assert download.status_code == 200
    prs = Presentation(io.BytesIO(download.data))
    assert len(prs.slides) == 1
    assert prs.slides[0].shapes[1].text == 'Editable text'


def test_endpoint_rejects_unknown_mode(client, sample_project):
    response = client.post(f"/api/projects/{sample_project['project_id']}/export/editable-pptx",
                           json={'export_mode': 'invalid'})
    assert response.status_code == 400
