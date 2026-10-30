"""Clustering preserves layout, semantic color differences, and export-mode scope."""
import zipfile
from lxml import etree

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from services.export_service import ExportService
from services.image_editability.data_models import BBox, EditableElement, EditableImage
from utils.text_style_clustering import consolidate_text_styles


def line(slide, text, x, y, size, color=(30, 30, 30), factor=1):
    shape = slide.shapes.add_textbox(Pt(x*factor), Pt(y*factor), Pt(180*factor), Pt(24*factor))
    shape.text = text
    p = shape.text_frame.paragraphs[0]
    p.font.size = Pt(size*factor)
    p.font.color.rgb = RGBColor(*color)
    return shape


def size(shape):
    return shape.text_frame.paragraphs[0].runs[0].font.size.pt


@pytest.mark.parametrize('factor', [1, 1.6])
def test_vertical_lines_share_median_without_merging_columns_or_title(factor):
    prs = Presentation()
    prs.slide_width = Inches(40/3*factor)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    a = line(slide, 'First line', 20, 60, 18, factor=factor)
    b = line(slide, 'Second line', 20, 90, 19, factor=factor)
    c = line(slide, 'Third line', 20, 120, 20, factor=factor)
    # The same height, but no horizontal overlap: independent column.
    other = line(slide, 'Other column', 240, 90, 19, factor=factor)
    title = line(slide, 'Title', 20, 30, 42, factor=factor)
    before = [(s.text,s.left,s.top,s.width,s.height) for s in slide.shapes]
    stats = consolidate_text_styles(slide, prs.slide_width)
    assert [size(s) for s in (a,b,c)] == pytest.approx([19*factor]*3, abs=.02)
    assert other.text_frame.paragraphs[0].font.size.pt == pytest.approx(19*factor)
    assert title.text_frame.paragraphs[0].font.size.pt == pytest.approx(42*factor)
    assert before == [(s.text,s.left,s.top,s.width,s.height) for s in slide.shapes]
    assert len(stats['font_clusters']) == 1
    assert stats['font_sizes_changed'] == 2


def test_color_medians_preserve_multicolor_runs_and_white_vs_tinted_text():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    a = line(slide,'Body',20,50,20,(30,30,30))
    b = line(slide,'More',20,80,20,(34,34,34))
    accent = b.text_frame.paragraphs[0].add_run()
    accent.text = ' highlighted'
    accent.font.size = Pt(20)
    accent.font.color.rgb = RGBColor(220,70,40)
    white = line(slide,'White',20,160,20,(250,250,250))
    green = line(slide,'Green',20,200,20,(244,250,244))
    consolidate_text_styles(slide, prs.slide_width)
    assert a.text_frame.paragraphs[0].runs[0].font.color.rgb == RGBColor(32,32,32)
    assert b.text_frame.paragraphs[0].runs[0].font.color.rgb == RGBColor(32,32,32)
    assert accent.font.color.rgb == RGBColor(220,70,40)
    assert white.text_frame.paragraphs[0].runs[0].font.color.rgb == RGBColor(250,250,250)
    assert green.text_frame.paragraphs[0].runs[0].font.color.rgb == RGBColor(244,250,244)


def test_small_heading_and_bold_text_do_not_merge_into_body():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading = line(slide, 'Heading', 20, 30, 14)
    body = line(slide, 'Body', 20, 60, 10)
    bold = line(slide, 'Emphasis', 20, 90, 11)
    bold.text_frame.paragraphs[0].font.bold = True
    stats = consolidate_text_styles(slide, prs.slide_width)
    assert stats['font_clusters'] == []
    assert heading.text_frame.paragraphs[0].font.size.pt == 14
    assert body.text_frame.paragraphs[0].font.size.pt == 10


def test_single_ocr_line_fits_width_and_stays_single_line():
    from utils.text_only_layout import fit_text_shape, measured_font, font_family
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = line(slide, '追求卓越的心', 20, 30, 30)
    shape.width = Pt(60)
    shape.height = Pt(40)
    fit_text_shape(shape, initialize=True)
    assert measured_font(size(shape)).getlength(shape.text) / 4 <= 60
    assert shape.text_frame.word_wrap is False
    assert shape.text_frame.paragraphs[0].runs[0].font.name == font_family()


def test_mixed_size_text_and_native_equations_are_not_font_clusters():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    a = line(slide,'Mixed',20,50,20)
    small = a.text_frame.paragraphs[0].add_run()
    small.text = ' footnote'
    small.font.size = Pt(10)
    line(slide,'Next',20,80,22)
    math_shape = slide.shapes.add_textbox(Pt(20),Pt(110),Pt(180),Pt(24))
    stats = consolidate_text_styles(slide, prs.slide_width)
    assert stats['font_clusters'] == []
    assert small.font.size.pt == 10
    assert math_shape.text == ''


@pytest.mark.parametrize('mode', ['simple','full'])
def test_export_applies_clustering_only_to_text_only_mode(tmp_path, monkeypatch, mode):
    from utils.pptx_builder import PPTXBuilder
    source = tmp_path/'source.png'
    background = tmp_path/'background.png'
    Image.new('RGB',(1280,720),'white').save(source)
    Image.new('RGB',(1280,720),'gray').save(background)
    monkeypatch.setattr(PPTXBuilder,'calculate_font_size',lambda self,bbox,text,*args: {'First':18,'Second':22}[text])
    elements = [EditableElement(element_id=f't{i}',element_type='text',
                bbox=BBox(20,60+i*40,260,84+i*40),
                bbox_global=BBox(20,60+i*40,260,84+i*40),content=text)
                for i,text in enumerate(['First','Second'])]
    img = EditableImage(image_id='test',image_path=str(source),width=1280,height=720,
                        clean_background=str(background),elements=elements)
    output=tmp_path/f'{mode}.pptx'
    ExportService.create_editable_pptx_with_recursive_analysis(
        editable_images=[img],export_mode=mode,output_file=str(output),
        slide_width_pixels=1280,slide_height_pixels=720,
    )
    shapes=[s for s in Presentation(str(output)).slides[0].shapes if s.has_text_frame]
    actual=[(s.text_frame.paragraphs[0].runs[0].font.size or s.text_frame.paragraphs[0].font.size).pt for s in shapes]
    if mode == 'simple':
        assert actual[0] == pytest.approx(actual[1], abs=.02)
        assert all(s.text_frame.word_wrap is False for s in shapes)
        with zipfile.ZipFile(output) as package:
            assert 'ppt/fonts/banana-text.fntdata' in package.namelist()
            presentation = etree.fromstring(package.read('ppt/presentation.xml'))
            assert presentation.get('embedTrueTypeFonts') == '1'
            assert presentation.get('saveSubsetFonts') == '0'
    else:
        assert actual == pytest.approx([18,22])
        with zipfile.ZipFile(output) as package:
            assert 'ppt/fonts/banana-text.fntdata' not in package.namelist()
