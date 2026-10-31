"""Deterministic OCR-line layout using the same bundled font for fit and delivery."""
from functools import lru_cache
from pathlib import Path

from PIL import ImageFont
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

FONT_PATH = Path(__file__).resolve().parent.parent / 'fonts' / 'NotoSansSC-Regular.ttf'


@lru_cache(maxsize=256)
def measured_font(size):
    return ImageFont.truetype(str(FONT_PATH), max(1, round(size * 4)))


def font_family():
    return measured_font(12).getname()[0]


def fit_text_shape(shape, *, initialize=False):
    """Preserve explicit OCR lines, never use automatic wrapping to fit a line."""
    frame = shape.text_frame
    frame.word_wrap = False
    frame.auto_size = None
    lines = shape.text.splitlines() or ['']
    sizes = []
    bold = False
    for paragraph in frame.paragraphs:
        if paragraph.font.size is not None:
            sizes.append(paragraph.font.size.pt)
        for run in paragraph.runs:
            if run.font.size is not None:
                sizes.append(run.font.size.pt)
            bold |= bool(run.font.bold if run.font.bold is not None else paragraph.font.bold)
    if not sizes:
        return
    width = shape.width / 12700
    height = shape.height / 12700
    # Renderers may add CJK/Latin separation. Reserving a quarter em at each
    # transition accounts for that without changing the recognized characters.
    def line_width(line, size):
        transitions = sum((ord(a) > 255) != (ord(b) > 255)
                          for a, b in zip(line, line[1:]) if not a.isspace() and not b.isspace())
        return measured_font(size).getlength(line) / 4 * (1.04 if bold else 1) + transitions * size * .25
    cap = min(height / len(lines), 200) if initialize else min(min(sizes), height / len(lines))
    low, high = .5, max(.5, cap)
    for _ in range(16):
        size = (low + high) / 2
        if max(line_width(line, size) for line in lines) <= width * .98:
            low = size
        else:
            high = size
    size = round(low, 2)
    family = font_family()
    for paragraph in frame.paragraphs:
        paragraph.font.size = Pt(size)
        paragraph.font.name = family
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.name = family
            properties = run._r.get_or_add_rPr()
            properties.set('kern', '0')
            for tag in ('a:ea', 'a:cs'):
                node = properties.find('{http://schemas.openxmlformats.org/drawingml/2006/main}' + tag[2:])
                if node is None:
                    node = OxmlElement(tag)
                    properties.append(node)
                node.set('typeface', family)
