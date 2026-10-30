"""Page-local text style clustering for text-only editable exports.

Ported from PPT Layer Restore's restore_text/consolidation.js and
gemini_style_consolidation.js. Font grouping follows its Edit-Banana adaptation:
nearby, horizontally overlapping boxes with similar sizes share a median size.
Apply after bbox fitting so the builder cannot overwrite the clustered sizes.
"""
import colorsys
import math
from statistics import median

from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_COLOR_TYPE
from pptx.util import Pt


def _rounded(value):
    # Match JavaScript Math.round for nonnegative sizes and RGB channels.
    return math.floor(value + 0.5)


def _median_color(colors):
    return tuple(_rounded(median(c[i] for c in colors)) for i in range(3))


def _color_family(rgb):
    r, g, b = rgb
    spread = max(rgb) - min(rgb)
    h, lightness, saturation = colorsys.rgb_to_hls(r / 255, g / 255, b / 255)
    h *= 360
    if min(rgb) >= 245 and spread <= 2:
        return 'neutral_white'
    if g - max(r, b) >= 3 or (
        g >= r and g >= b and 70 <= h <= 175
        and (spread >= 4 or saturation >= .04)
    ):
        return 'green'
    if min(r, g) - b >= 8 or (35 <= h < 70 and saturation >= .08):
        return 'yellow'
    if spread <= 10 and saturation < .08:
        return 'neutral_white' if lightness > .9 else 'neutral_gray'
    if h < 25 or h >= 345:
        return 'red'
    for limit, family in [(35, 'orange'), (70, 'yellow'), (175, 'green'),
                          (205, 'cyan'), (260, 'blue'), (345, 'purple')]:
        if h < limit:
            return family
    return 'unknown'


def _neighbors(a, b):
    sa, sb = a['shape'], b['shape']
    gap_ab = sb.top - (sa.top + sa.height)
    gap_ba = sa.top - (sb.top + sb.height)
    gap = 0 if gap_ab < 0 and gap_ba < 0 else min(abs(gap_ab), abs(gap_ba))
    overlap = not (sa.left + sa.width < sb.left or sb.left + sb.width < sa.left)
    close_size = max(a['size'], b['size']) / min(a['size'], b['size']) <= 1.16
    same_style = a['bold'] == b['bold'] and a['family'] == b['family']
    return gap < min(sa.height, sb.height) * 1.25 and overlap and close_size and same_style


def consolidate_text_styles(slide, slide_width):
    """Unify fitted text styles without moving boxes, joining lines, or changing text.

    Group only similar sizes (at most 16% apart), weights, and font families.
    Color clustering is page-wide, within color families and an
    RGB distance of 34. Each colored run remains a separate color sample, so
    intentional multicolor text is not flattened to the paragraph's base color.
    """
    scale = 960 / (slide_width / 12700)
    nodes, colors = [], []
    for shape in sorted(slide.shapes, key=lambda s: (s.top, s.left)):
        if not shape.has_text_frame or not shape.text.strip() or shape.width <= 0 or shape.height <= 0:
            continue
        sizes, fonts, weights, families = [], [], [], []
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if not run.text.strip():
                    continue
                size = run.font.size if run.font.size is not None else paragraph.font.size
                if size is not None:
                    sizes.append(size.pt)
                    fonts.append(run.font)
                    weights.append(bool(run.font.bold if run.font.bold is not None else paragraph.font.bold))
                    families.append(run.font.name or paragraph.font.name)
                color = run.font.color
                if color.type is None:
                    color = paragraph.font.color
                if color.type == MSO_COLOR_TYPE.RGB:
                    colors.append({'font': run.font, 'rgb': tuple(color.rgb), 'id': shape.shape_id})
        # Native equations and deliberately mixed-size boxes are not line samples.
        if sizes and max(sizes) - min(sizes) < .01:
            nodes.append({'shape': shape, 'fonts': fonts, 'size': sizes[0] * scale,
                          'bold': tuple(sorted(set(weights))), 'family': tuple(sorted(set(families), key=str))})

    parent = list(range(len(nodes)))

    def find(index):
        while parent[index] != index:
            parent[index] = parent[parent[index]]
            index = parent[index]
        return index

    for i, a in enumerate(nodes):
        for j in range(i + 1, len(nodes)):
            if _neighbors(a, nodes[j]):
                members = [n['size'] for k, n in enumerate(nodes) if find(k) in (find(i), find(j))]
                # Do not let a chain of small pairwise differences merge a heading
                # with body copy several lines below it.
                if max(members) / min(members) <= 1.16:
                    parent[find(i)] = find(j)
    groups = {}
    for i, node in enumerate(nodes):
        groups.setdefault(find(i), []).append(node)

    diagnostics = {'font_clusters': [], 'font_sizes_changed': 0,
                   'color_clusters': [], 'colors_changed': 0}
    for group in groups.values():
        if len(group) < 2:
            continue
        size = _rounded(median(n['size'] for n in group) * 10) / 10 / scale
        members = []
        for node in group:
            before = node['size'] / scale
            for font in node['fonts']:
                font.size = Pt(size)
            for paragraph in node['shape'].text_frame.paragraphs:
                paragraph.font.size = Pt(size)
            diagnostics['font_sizes_changed'] += abs(before - size) > .1
            members.append({'shape_id': node['shape'].shape_id, 'before': before, 'after': size})
        diagnostics['font_clusters'].append(members)

    clusters = []
    for item in colors:
        family = _color_family(item['rgb'])
        candidates = [c for c in clusters if c['family'] == family]
        cluster = min(candidates, key=lambda c: math.dist(item['rgb'], c['center']), default=None)
        if cluster is None or math.dist(item['rgb'], cluster['center']) > 34:
            cluster = {'family': family, 'center': item['rgb'], 'items': []}
            clusters.append(cluster)
        cluster['items'].append(item)
        cluster['center'] = _median_color([m['rgb'] for m in cluster['items']])
    for cluster in clusters:
        for item in cluster['items']:
            item['font'].color.rgb = RGBColor(*cluster['center'])
            diagnostics['colors_changed'] += item['rgb'] != cluster['center']
        diagnostics['color_clusters'].append({
            'family': cluster['family'], 'color': cluster['center'],
            'shape_ids': [item['id'] for item in cluster['items']],
        })
    return diagnostics
